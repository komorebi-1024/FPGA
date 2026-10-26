"""Build the beginner manual and navigation artifacts from reviewed Markdown."""
from pathlib import Path
import csv,json,re,shutil,hashlib,wave,math,struct
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
root=Path(__file__).resolve().parent.parent;docs=root/'00_Docs';sim=root/'04_Sim'
audit=json.loads((root/'06_Delivery/release_audit.json').read_text('utf-8'))
# Reproducible visual sheet made solely from the RTL simulator's output files.
sheet=Image.new('RGB',(1280,1000),(245,247,249));draw=ImageDraw.Draw(sheet)
f=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',22)
draw.text((14,7),'四站 RTL 仿真输出  非实机照片',font=f,fill=(22,35,48))
for i in range(4):sheet.paste(Image.open(sim/f'osd_rtl_{i}.png'),((i%2)*640,40+(i//2)*480))
sheet.save(sim/'四站RTL仿真输出.png')
entries=list(csv.DictReader((root/'02_Examples/00_工程入口.csv').open(encoding='utf-8-sig')))
explanations={
'0_cnt1s':'一秒计数与时钟分频，认识时钟节拍。',
'1_breath_led':'呼吸灯与 PWM，观察占空比如何改变亮度。',
'2_key_led':'按键控制 LED，认识输入输出和引脚约束。',
'3_seg_LDE7S':'七段数码管译码，理解数字到段码的映射。',
'4_seg_led_static_top':'数码管静态显示，学习段选与位选。',
'5_seg_SCAN':'数码管动态扫描，多位共用段线。',
'6_cnt10':'十进制计数与数码显示，学习复位和进位。',
'7_seg_led':'数码显示综合例程，对照显示顶层的模块连接。',
'8_key_sw':'按键和拨码开关输入，适合核对实体控制。',
'9_MUX_buzz':'选择逻辑与蜂鸣器输出；不是 HDMI 音频。',
'10_PLL_TEST':'PLL 时钟生成，观察倍频与分频配置。',
'11_Di_TRemble':'按键消抖，理解一次按下为何产生多个电平跳变。',
'12_MUL_4B':'乘法器实验；目录与顶层位数命名不同，以源码为准。',
'13_music':'板上音乐或蜂鸣器播放，认识音调与节拍。',
'14_ps2_piano':'PS2 键盘控制电子琴，需要对应外设。',
'15_VGA':'VGA 显示时序，与本项目行场消隐概念相关。',
'16_uart_rs232':'串口收发测试，学习串行数据与波特率。',
'17_DC_MOTO':'直流电机控制，需要驱动模块与独立负载。',
'18_sd_card_audio':'SD 卡音频读取，适合了解文件数据与播放节奏。',
'19_ball_top':'VGA 小球显示或游戏，关注位置状态与画面更新。',
'20_i2sound_wm8731':'I2S 与 WM8731 音频芯片接口，需要对应音频硬件。',
'21_liushuideng':'流水灯，学习计数器驱动状态变化。',
'22_jiaotong':'交通灯状态机，适合初学状态迁移。',
'23_SCAN_4X4KEY':'4×4 矩阵键盘扫描，需要键盘扩展连接。',
'24_VGA_shiyan9_B':'VGA 配套实验工程，结合实验文档查看具体画面。',
'25_seg70_small16x16':'16×16 点阵或显示组合实验，结合文档核对接口。',
'26_LCD2004_A':'字符 LCD 实验；顶层名 lcd1602，与目录命名不同。',
'27_DAC_SIN':'12 位 DAC 正弦输出，需要外部 DAC 模块。',
'28_adda_direct':'ADC/DAC 数据直通，学习采样和转换数据通路。',
'29_DA2_10B':'双路 10 位 DAC 相关实验，需要对应转换模块。',
'30_LCD_ATK':'ATK LCD 显示接口实验，不能直接代替 HDMI。',
'31_hdmi_tx_display_A':'HDMI 显示参考 A，学习视频发送及引脚。',
'31_hdmi_tx_display_B':'HDMI 显示参考 B，对照 A 的配置和连接差异。',
'32_Ethernet_test':'以太网 UDP 发送示例，需要网口环境。',
'33_HX4S20_Etherent_test_01_02_final_1000':'HX4S20 千兆以太网测试，与本次基础播放器不直接相关。',
'34_GPS':'GPS 接收/显示示例，仅供后续定位扩展研究。',
'HX_CESHI_EG4S_BORD':'板卡综合测试，配套外设和接线须逐项核对。'
}
rows=[];full=[]
for e in entries:
 p=Path(e['工程入口']);parts=p.parts;group=e['来源分组']
 if group=='02_DEMO':
  name=parts[3];desc=explanations[name];label='DEMO'
 elif group=='01_Contest':
  name=parts[2];label='比赛'
  desc='TF 图片与 HDMI 显示参考，目录内 .al 名称带 i2s，注意以路径区分。' if name=='lab_ex4_tf' else 'TF 图像与 HDMI 音频综合参考，本项目的原始基线。'
 elif group=='03_IP_Reference':
  name=parts[2].split('_')[0];label='IP'
  desc={'APUG010':'SDRAM 本地总线读写参考，了解内存接口。','APUG011':'高性能 SDRAM 读写参考，了解缓存后端。','APUG092':'HDMI 1.4b 发射器参考，包含音视频复用接口。'}[name]
 else:name='Serial_Display_Dual';label='上位机';desc='双通道串口波形显示配套工程，需参考其通信格式。'
 rows.append((label,name,desc));full.append({**e,'名称':name,'用途':desc})
with (root/'02_Examples/01_带用途的工程索引.csv').open('w',encoding='utf-8-sig',newline='') as fcsv:
 w=csv.DictWriter(fcsv,fieldnames=list(full[0]));w.writeheader();w.writerows(full)
index=['# 原始例程入口与用途','',f'共 {len(rows)} 个 .al 入口；只重新编译了比赛音视频基线，其余未逐个验证。','']
for i,e in enumerate(full,1):index += [f'## {i} {e["名称"]}',e['用途'],f'入口：{e["工程入口"]}',f'顶层：{e["顶层"]}；状态：{e["验证状态"]}','']
(root/'02_Examples/01_例程说明.md').write_text('\n'.join(index),'utf-8')
table='|分组|原目录或参考编号|用途与学习重点|\n|---|---|---|\n'+'\n'.join('|'+ '|'.join(r)+'|' for r in rows)
mdpath=docs/'随行视听西湖工程上手与复现教程.md'
md=mdpath.read_text('utf-8')
if '{{EXAMPLE_TABLE}}' in md:md=md.replace('{{EXAMPLE_TABLE}}',table);mdpath.write_text(md,'utf-8')
# Create compact workflow launchers; no launcher performs physical-disk writes.
for filename,action in [('01_Run_Simulation.cmd','Sim'),('02_Build_Project.cmd','Build'),('03_Rebuild_Assets.cmd','Assets')]:
 (root/filename).write_bytes(('@echo off\r\npowershell.exe -NoProfile -ExecutionPolicy Bypass -File "%~dp0tools\\workflow.ps1" -Action '+action+'\r\nif errorlevel 1 echo FAILED: inspect the log.\r\npause\r\n').encode('ascii'))
readme='''# 从这里开始

本目录是随行视听西湖工程。先读 00_Docs 中的《随行视听西湖工程上手与复现教程.docx》。

- 四张实景原图：01_Assets/Original；标准图片：01_Assets/TF_BMP。
- 所有例程：02_Examples；43 个工程入口见带用途的工程索引。
- 正式工程：03_Project/WestLake_AV/src/td_project/HDMI1.4b_Transmitter_v1.0.al。
- 双击 01_Run_Simulation.cmd 跑自检；双击 02_Build_Project.cmd 重新编译。
- 当前下载文件：05_Build/westlake/westlake.bit。
- TF 卡使用 01_Assets/WestLake_TF_128MiB.img 整盘镜像，不能只拖入 BMP。
- 改字幕后运行 03_Rebuild_Assets.cmd，再仿真和编译。

已完成：素材、源代码、七组自检仿真、TD 布局布线和受约束路径时序检查。
待下午实测：TF 卡及 HDMI 兼容性、声画质量、持续运行、Flash 固化与冷启动。
任何模拟图都不是实机截图。原始资料没有修改。根目录按钮不会写入物理磁盘。
'''
(root/'README_从这里开始.md').write_text(readme,'utf-8')
state=f'''# 工作状态

2026-10-25 软件交付。用户已授权整理素材、例程、工程、仿真、编译与教程。
硬件在队友处；当前未上板，不得声称硬件通过。

四站：少年宫（断桥）→涌金门→雷峰塔→岳庙（跨虹桥）。
完成：43 例程索引、四张授权照片、4 个标准 BMP、128 MiB FAT32 镜像、7 个新 RTL 模块、7 组自检仿真、TD 编译与位流。
时序：SWNS {audit['setup_wns_ns']} ns，HWNS {audit['hold_wns_ns']} ns，STA 覆盖 {audit['sta_coverage_percent']}%。
仍有原厂 IP 告警，外部 IO 延时未完整约束；板上验证待做。
正式下载：05_Build/westlake/westlake.bit
SHA256：{audit['bit_sha256']}

下一步：阅读教程并亲自运行入口；下午按上板验收表完成 TF、HDMI、声音和冷启动测试。
'''
(docs/'工作状态.md').write_text(state,'utf-8')
checks=['首图启动','四站顺序与字幕','按键短按长按','自动播放四档','切换无撕裂','HDMI持续测试音','到站提示音','动态时间与模式','10分钟连续运行','无卡与坏镜像','Flash型号与校验','三次断电启动']
with (docs/'上板验收表.csv').open('w',encoding='utf-8-sig',newline='') as fcsv:
 w=csv.writer(fcsv);w.writerow(['测试项','状态','实际结果','设备型号','日期与测试人','证据文件','使用位流SHA256'])
 for c in checks:w.writerow([c,'待实测','','','','',audit['bit_sha256']])

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.72);section.bottom_margin=Inches(.72);section.left_margin=Inches(.72);section.right_margin=Inches(.72)
section.header_distance=Inches(.3);section.footer_distance=Inches(.3)
def fontstyle(style,size,bold=False):
 style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 rfonts=style.element.get_or_add_rPr().rFonts
 for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rfonts.attrib.pop(qn('w:'+key),None)
 for key in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+key),'Microsoft YaHei')
for name,size in [('Normal',11),('Title',23),('Subtitle',13),('Heading 1',17),('Heading 2',13),('Heading 3',11.5)]:fontstyle(doc.styles[name],size,name.startswith('Heading'))
# The linked character styles generated by the Word template can otherwise make
# WPS map a Chinese Title paragraph to MS Gothic. Keep the linked styles explicit.
for style_el in doc.styles.element.iter(qn('w:style')):
 if style_el.get(qn('w:styleId')) in {'Title','TitleChar','Heading1','Heading1Char','Heading2','Heading2Char','Heading3','Heading3Char','Normal'}:
  rpr=style_el.find(qn('w:rPr'))
  if rpr is None:rpr=OxmlElement('w:rPr');style_el.append(rpr)
  rfonts=rpr.find(qn('w:rFonts'))
  if rfonts is None:rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rfonts.attrib.pop(qn('w:'+key),None)
  for key in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+key),'Microsoft YaHei')
normal=doc.styles['Normal'].paragraph_format;normal.line_spacing=1.15;normal.space_after=Pt(6)
for hn in ['Heading 1','Heading 2','Heading 3']:
 st=doc.styles[hn].paragraph_format;st.space_before=Pt(14);st.space_after=Pt(7);st.keep_with_next=True
code=doc.styles.add_style('WLCode',1);fontstyle(code,9);code.font.name='Consolas';code.paragraph_format.line_spacing=1.1;code.paragraph_format.space_after=Pt(2)
header=section.header.paragraphs[0];header.text='随行视听西湖工程  上手与复现';header.runs[0].font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('WestLake AV  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9)
def rich(p,text):
 for i,part in enumerate(re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text)):
  if not part:continue
  if part.startswith('`') and part.endswith('`'):
   r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(10)
  elif part.startswith('**'):r=p.add_run(part[2:-2]);r.bold=True
  else:p.add_run(part)
def make_table(lines):
 records=[[c.strip() for c in ln.strip().strip('|').split('|')] for ln in lines]
 records=[r for r in records if not all(re.fullmatch(r':?-+:?',c) for c in r)]
 cols=len(records[0]);t=doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=([1.48,2.12,3.42] if cols==3 else ([2.2,4.82] if cols==2 else [7.02/cols]*cols))
 widths_by_header={'项目':[1.25,3.15,2.62],'顺序':[.5,1.8,4.72],'目录':[2.3,2.8,1.92],
 '术语':[1.1,2.7,3.22],'文件':[1.85,2.1,3.07],'数码管状态':[1.1,2.9,3.02],
 '实体控制':[1.1,2.0,3.92],'测试台':[1.6,3.2,2.22],'指标':[2.1,1.6,3.32],
 '现象':[1.55,2.65,2.82],'用时':[1.45,2.9,2.67],'测试项目':[1.6,3.8,1.62]}
 widths=widths_by_header.get(records[0][0],widths)
 if records[0][0]=='分组':widths=[.57,2.35,4.10]
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 pr.append(borders)
 for j,record in enumerate(records):
  row=t.add_row();trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  if j==0:trpr.append(OxmlElement('w:tblHeader'))
  for k,(cell,txt) in enumerate(zip(row.cells,record)):
   cell.width=Inches(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');marg.append(e)
   tcpr.append(marg)
   if j==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EFF5');tcpr.append(shade)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.12
   # Keep the final two rows together instead of leaving one orphaned row.
   if j==len(records)-2:p.paragraph_format.keep_with_next=True
   rich(p,txt)
   for r in p.runs:r.font.size=Pt(10);r.bold=j==0
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
lines=md.splitlines();i=0;in_code=False
while i<len(lines):
 line=lines[i]
 if line.startswith('```'):in_code=not in_code;i+=1;continue
 if in_code:
  p=doc.add_paragraph(line,style='WLCode');p.paragraph_format.keep_with_next=i+1<len(lines) and not lines[i+1].startswith('```');i+=1;continue
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  tablelines=[]
  while i<len(lines) and lines[i].startswith('|'):tablelines.append(lines[i]);i+=1
  make_table(tablelines);continue
 if line.startswith('!['):
  match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str((docs/match[2]).resolve()),width=Inches(6.6))
 elif line.startswith('# '):
  title=doc.add_paragraph(line[2:],style='Title')
  for run in title.runs:
   run.font.name='Microsoft YaHei'
   rfonts=run._element.get_or_add_rPr().rFonts
   for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rfonts.attrib.pop(qn('w:'+key),None)
   for key in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+key),'Microsoft YaHei')
 elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
 elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
 else:
  p=doc.add_paragraph();rich(p,line)
  if line.startswith('1. ') and i+1<len(lines) and lines[i+1].startswith('2. '):p.paragraph_format.keep_with_next=True
 i+=1
doc.core_properties.title='随行视听西湖工程上手与复现教程';doc.core_properties.subject='HX4S20C FPGA 西湖四站音视频工程'
doc.core_properties.author='项目辅助文档';doc.core_properties.keywords='FPGA;Anlogic;HX4S20C;WestLake'
# The bundled default Word template can carry a blue paragraph border on Title.
# Remove paragraph borders from all styles and body paragraphs explicitly.
for tree in [doc.styles.element,doc.element]:
 for element in list(tree.iter(qn('w:pBdr'))):element.getparent().remove(element)
output=docs/'随行视听西湖工程上手与复现教程.docx';doc.save(output)
print(f'Created {output}; {len(rows)} example entries; {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.')
